"""Read the file types LocalRAG supports."""

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = frozenset({".docx", ".md", ".pdf", ".txt"})


class IngestionError(ValueError):
    """Raised when a document cannot be read."""


@dataclass(frozen=True, slots=True)
class Document:
    """A document and its path."""

    path: Path
    content: str


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as error:
        raise IngestionError(
            "PDF support requires the optional document dependencies; "
            "install with: pip install 'localrag[documents]'"
        ) from error

    try:
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (OSError, PdfReadError, ValueError) as error:
        raise IngestionError(f"Could not read PDF document: {path}") from error


def _read_docx(path: Path) -> str:
    try:
        from docx import Document as WordDocument
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as error:
        raise IngestionError(
            "Word support requires the optional document dependencies; "
            "install with: pip install 'localrag[documents]'"
        ) from error

    try:
        word_document = WordDocument(str(path))
        paragraphs = [paragraph.text for paragraph in word_document.paragraphs]
        tables = [
            " | ".join(cell.text for cell in row.cells)
            for table in word_document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + tables)
    except (OSError, PackageNotFoundError, ValueError) as error:
        raise IngestionError(f"Could not read Word document: {path}") from error


def read_document(path: Path) -> Document:
    """Read one supported document."""
    path = path.expanduser()

    if not path.is_file():
        raise IngestionError(f"Document does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise IngestionError(f"Unsupported document type: {path.suffix or '(none)'}; use {supported}")

    if suffix == ".pdf":
        content = _read_pdf(path)
    elif suffix == ".docx":
        content = _read_docx(path)
    else:
        try:
            content = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as error:
            raise IngestionError(f"Document is not valid UTF-8: {path}") from error
        except OSError as error:
            raise IngestionError(f"Could not read document: {path}") from error

    if not content.strip():
        raise IngestionError(f"Document is empty: {path}")

    return Document(path=path.resolve(), content=content)


def ingest_directory(directory: Path) -> list[Document]:
    """Read supported documents under a folder."""
    directory = directory.expanduser()

    if not directory.is_dir():
        raise IngestionError(f"Document directory does not exist: {directory}")

    documents: list[Document] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            documents.append(read_document(path))

    return documents
